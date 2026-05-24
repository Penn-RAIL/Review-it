import html
import os
import re
from datetime import datetime
from pathlib import Path
from uuid import uuid4

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.table import _Cell

from .ollama import generate_with_ollama
from .review_prompt import build_review_prompt


REPORT_DIR = Path(os.environ.get("REVIEW_IT_REPORT_DIR", Path.home() / ".review-it" / "reports"))

KNOWN_TABLE_HEADERS = [
    ["Location", "Original Text", "Issue Type", "Problem", "Suggested Revision"],
    ["Location", "Stated Value or Claim", "Potential Issue", "Why This May Be Incorrect or Unclear", "Recommended Verification or Correction"],
    ["Table/Figure", "Purpose", "Clarity", "Problems Identified", "Recommended Fix"],
    ["Location", "Claim", "Citation Issue", "Recommendation"],
    ["Location", "Original Claim", "Concern", "Suggested Softer or More Accurate Wording"],
    ["Priority", "Category", "Issue", "Recommended Action", "Estimated Effort"],
]

TABLE_STOP_PATTERNS = [
    r"^after the table",
    r"^then provide",
    r"^rules:?$",
    r"^overall language assessment:?$",
    r"^most common language issues:?$",
    r"^sections needing",
    r"^high-risk numerical issues:?$",
    r"^claims that need",
    r"^statements that should",
    r"^missing tables",
    r"^tables or figures",
    r"^caption issues:?$",
    r"^areas needing",
    r"^potentially missing",
    r"^most concerning",
    r"^likely reviewer concerns:?$",
    r"^potential reasons",
    r"^what would most improve",
    r"^recommended target journal",
    r"^rationale:?$",
    r"^major comments:?$",
    r"^minor comments:?$",
    r"^include at least:?$",
    r"^final recommendation:?$",
    r"^justification:?$",
    r"^top 5 changes",
    r"^additional instructions",
]

MAJOR_SECTION_HEADINGS = {
    "AI-Assisted Pre-Submission Manuscript Review Report",
    "Manuscript Readiness Summary",
    "Grammar and Language Check",
    "Text Organization and Flow",
    "Factual, Numerical, and Calculation Error Check",
    "Technical and Methodological Error Check",
    "Table and Figure Review",
    "Citation and Literature Support Check",
    "Overclaiming and Interpretation Check",
    "Journal Readiness and Reviewer Risk",
    "Reviewer-Style Comments",
    "Prioritized Revision Checklist",
    "Final Recommendation",
    "Additional Instructions",
}

SUBSECTION_HEADINGS = {
    "Title",
    "Abstract",
    "Introduction",
    "Methods",
    "Results",
    "Discussion",
    "Conclusion",
    "Study Design",
    "Data and Dataset Quality",
    "Model, Algorithm, or Intervention Details",
    "Statistical Analysis",
    "Evaluation Metrics",
    "Reproducibility and Transparency",
    "Ethics, Bias, Privacy, and Generalizability",
}


async def create_review_report(manuscript: dict, model: str, model_score: dict, system_profile: dict) -> dict:
    started_at = datetime.now()
    prompt = build_review_prompt(manuscript["text"])
    raw_markdown = await generate_with_ollama(
        model=model,
        prompt=prompt,
        options={"num_ctx": min(model_score.get("contextWindow", 8192), 131072)},
    )
    markdown = normalize_markdown(prepare_report_markdown(raw_markdown))
    log_llm_output(raw_markdown, markdown)

    report = {
        "id": str(uuid4()),
        "model": model,
        "createdAt": started_at.isoformat(),
        "markdown": markdown,
        "rawMarkdown": raw_markdown,
        "html": markdown_to_html(markdown),
        "metadata": {
            "manuscript": manuscript["profile"],
            "modelScore": model_score,
            "systemProfile": system_profile,
            "promptMode": "fixed pre-submission peer review prompt",
        },
    }
    report["docxPath"] = str(write_docx(report))
    return report


def log_llm_output(raw_markdown: str, normalized_markdown: str) -> None:
    log_dir = Path(os.environ.get("REVIEW_IT_LOG_DIR", Path.home() / ".review-it" / "logs"))
    log_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
    (log_dir / f"{stamp}-raw.md").write_text(raw_markdown, encoding="utf-8")
    (log_dir / f"{stamp}-normalized.md").write_text(normalized_markdown, encoding="utf-8")
    print(f"Review It LLM output logged to {log_dir / f'{stamp}-raw.md'}", flush=True)


def write_docx(report: dict) -> Path:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    file_path = REPORT_DIR / f"{report['id']}.docx"

    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(11)

    title = document.add_heading("Manuscript Review Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(report["metadata"]["manuscript"]["fileName"]).bold = True
    subtitle.add_run(f" | {datetime.fromisoformat(report['createdAt']).strftime('%Y-%m-%d %H:%M')}")

    add_markdown(document, report["markdown"])
    document.save(file_path)
    return file_path


def add_markdown(document: Document, markdown: str) -> None:
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        known_header = known_table_header(lines[index])
        if known_header:
            rows, index = collect_loose_table(lines, index, known_header)
            add_docx_table(document, rows)
            continue

        if is_markdown_table(lines, index):
            rows, index = collect_markdown_table(lines, index)
            add_docx_table(document, rows)
            continue

        if is_tab_table(lines, index):
            rows, index = collect_tab_table(lines, index)
            add_docx_table(document, rows)
            continue

        raw_line = lines[index]
        line = raw_line.strip()
        if not line:
            document.add_paragraph("")
            index += 1
            continue
        if line.startswith("# "):
            document.add_heading(line[2:].strip(), level=1)
            index += 1
            continue
        if line.startswith("## "):
            document.add_heading(line[3:].strip(), level=2)
            index += 1
            continue
        if line.startswith("### "):
            document.add_heading(line[4:].strip(), level=3)
            index += 1
            continue
        if is_bold_heading(line):
            document.add_heading(clean_inline_markdown(line), level=1 if is_major_heading(line) else 2)
            index += 1
            continue
        if is_numbered_subsection_heading(line):
            document.add_heading(clean_inline_markdown(re.sub(r"^\d+\.\d+\s+", "", line)), level=2)
            index += 1
            continue
        if is_numbered_major_heading(line):
            document.add_heading(clean_inline_markdown(re.sub(r"^\d+\.\s+", "", line)), level=1)
            index += 1
            continue
        if re.match(r"^[-*]\s+", line):
            add_text_paragraph(document, re.sub(r"^[-*]\s+", "", line), style="List Bullet")
            index += 1
            continue
        if re.match(r"^\d+\.\s+", line):
            add_text_paragraph(document, re.sub(r"^\d+\.\s+", "", line), style="List Number")
            index += 1
            continue
        add_text_paragraph(document, line)
        index += 1


def prepare_report_markdown(markdown: str) -> str:
    lines = markdown.splitlines()
    output = []
    index = 0

    while index < len(lines):
        known_header = known_table_header(lines[index])
        if known_header:
            rows, index = collect_loose_table(lines, index, known_header)
            output.extend(rows_to_markdown_table(rows))
            output.append("")
            continue

        if is_markdown_table(lines, index):
            rows, index = collect_markdown_table(lines, index)
            output.extend(rows_to_markdown_table(rows))
            output.append("")
            continue

        if is_tab_table(lines, index):
            rows, index = collect_tab_table(lines, index)
            output.extend(rows_to_markdown_table(rows))
            output.append("")
            continue

        line = lines[index].strip()
        if is_bold_heading(line):
            level = "#" if is_major_heading(line) else "##"
            output.append(f"{level} {clean_inline_markdown(line)}")
        else:
            output.append(clean_inline_markdown(line) if line else "")
        index += 1

    return "\n".join(output)


def rows_to_markdown_table(rows: list[list[str]]) -> list[str]:
    rows = normalize_table_rows(rows)
    if not rows:
        return []
    header = "| " + " | ".join(clean_table_cell(cell) for cell in rows[0]) + " |"
    separator = "| " + " | ".join("---" for _ in rows[0]) + " |"
    body = ["| " + " | ".join(clean_table_cell(cell) for cell in row) + " |" for row in rows[1:]]
    return [header, separator, *body]


def clean_table_cell(text: str) -> str:
    return clean_inline_markdown(text).replace("|", "/").strip()


def add_docx_table(document: Document, rows: list[list[str]]) -> None:
    rows = normalize_table_rows(rows)
    if not rows or not rows[0]:
        return

    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, rows[0]):
        write_cell(cell, value, bold=True)

    for row in rows[1:]:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            write_cell(cell, value)

    document.add_paragraph("")


def write_cell(cell: _Cell, value: str, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.text = ""
    run = paragraph.add_run(clean_inline_markdown(value))
    run.bold = bold


def add_text_paragraph(document: Document, text: str, style: str | None = None) -> None:
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(clean_inline_markdown(text))


def is_markdown_table(lines: list[str], index: int) -> bool:
    if index + 1 >= len(lines):
        return False
    current = lines[index].strip()
    separator = lines[index + 1].strip()
    return current.startswith("|") and "|" in current[1:] and is_markdown_separator(separator)


def is_markdown_separator(line: str) -> bool:
    if not line.startswith("|"):
        return False
    cells = parse_markdown_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def collect_markdown_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows = [parse_markdown_row(lines[index])]
    index += 2
    while index < len(lines) and lines[index].strip().startswith("|"):
        if not is_markdown_separator(lines[index].strip()):
            rows.append(parse_markdown_row(lines[index]))
        index += 1
    return rows, index


def parse_markdown_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_tab_table(lines: list[str], index: int) -> bool:
    line = lines[index].strip()
    return "\t" in line and len([cell for cell in line.split("\t") if cell.strip()]) > 1


def collect_tab_table(lines: list[str], index: int) -> tuple[list[list[str]], int]:
    rows = []
    while index < len(lines) and lines[index].strip() and "\t" in lines[index]:
        rows.append([cell.strip() for cell in lines[index].strip().split("\t")])
        index += 1
    return rows, index


def known_table_header(line: str) -> list[str] | None:
    stripped = line.strip()
    if stripped.startswith("|") and "|" in stripped[1:]:
        cells = parse_markdown_row(stripped)
        for header in KNOWN_TABLE_HEADERS:
            if [normalize_for_matching(cell) for cell in cells] == [normalize_for_matching(cell) for cell in header]:
                return header

    normalized = normalize_for_matching(line)
    for header in KNOWN_TABLE_HEADERS:
        if normalized == normalize_for_matching("\t".join(header)):
            return header
    return None


def collect_loose_table(lines: list[str], index: int, header: list[str]) -> tuple[list[list[str]], int]:
    block = []
    index += 1

    while index < len(lines):
        line = lines[index].strip()
        cleaned = clean_inline_markdown(line)

        if not line:
            index += 1
            continue

        if is_markdown_separator(line):
            index += 1
            continue

        if should_stop_table(cleaned, [header] + block):
            break

        if known_table_header(cleaned) and normalize_for_matching(cleaned) != normalize_for_matching("\t".join(header)):
            break

        block.append(cleaned)
        index += 1

    return [header] + parse_loose_table_block(header, block), index


def parse_loose_table_block(header: list[str], block: list[str]) -> list[list[str]]:
    block = [line for line in block if not is_example_row(line) and not is_markdown_separator(line)]
    if not block:
        return []

    explicit_rows = []
    for line in block:
        cells = split_table_line(line, len(header))
        if len(cells) >= len(header):
            explicit_rows.append(pad_row(cells[: len(header)], len(header)))
    if explicit_rows and len(explicit_rows) >= max(1, len(block) // 2):
        return explicit_rows

    if header[0] == "Location":
        groups = group_by_starts(block, looks_like_location_start)
        return [parse_location_group(header, group) for group in groups]

    if header[0] == "Table/Figure":
        groups = group_by_starts(block, looks_like_table_figure_start)
        return [parse_table_figure_group(group) for group in groups]

    if header[0] == "Priority":
        groups = group_by_starts(block, looks_like_priority_start)
        return [parse_priority_group(group) for group in groups]

    return [pad_row([line], len(header)) for line in block]


def is_example_row(line: str) -> bool:
    normalized = normalize_for_matching(line)
    return normalized.startswith("section/paragraph") or normalized.startswith("section/table/figure") or normalized.startswith("table 1 / figure 1") or normalized.startswith("high / medium / low")


def group_by_starts(block: list[str], start_predicate) -> list[list[str]]:
    groups = []
    current = []
    for line in block:
        if start_predicate(line) and current:
            groups.append(current)
            current = [line]
        else:
            current.append(line)
    if current:
        groups.append(current)
    return groups


def parse_location_group(header: list[str], group: list[str]) -> list[str]:
    location, first_value = split_location_prefix(group[0])
    rest = group[1:]

    if header == KNOWN_TABLE_HEADERS[0]:
        issue_type = infer_issue_type(group)
        problem = rest[-2] if len(rest) >= 2 else ""
        suggested = rest[-1] if rest else ""
        original_parts = [first_value] + rest[: max(0, len(rest) - 2)]
        return pad_row([location, " ".join(part for part in original_parts if part), issue_type, problem, suggested], len(header))

    if header == KNOWN_TABLE_HEADERS[1]:
        potential = rest[0] if len(rest) >= 1 else ""
        why = rest[1] if len(rest) >= 2 else ""
        correction = " ".join(rest[2:]) if len(rest) >= 3 else ""
        return pad_row([location, first_value, potential, why, correction], len(header))

    if header == KNOWN_TABLE_HEADERS[3]:
        issue = rest[0] if len(rest) >= 1 else infer_citation_issue(first_value)
        recommendation = " ".join(rest[1:]) if len(rest) >= 2 else ""
        return pad_row([location, first_value, issue, recommendation], len(header))

    if header == KNOWN_TABLE_HEADERS[4]:
        concern = rest[0] if len(rest) >= 1 else ""
        wording = " ".join(rest[1:]) if len(rest) >= 2 else ""
        return pad_row([location, first_value, concern, wording], len(header))

    return pad_row([location, first_value] + rest, len(header))


def parse_table_figure_group(group: list[str]) -> list[str]:
    item, purpose = split_table_figure_prefix(group[0])
    clarity = infer_clarity(group)
    problems = group[1] if len(group) >= 2 else ""
    fix = " ".join(group[2:]) if len(group) >= 3 else ""
    return pad_row([item, purpose, clarity, problems, fix], 5)


def parse_priority_group(group: list[str]) -> list[str]:
    priority_match = re.match(r"^(High|Medium|Low)\b\s*(.*)", group[0], re.I)
    priority = priority_match.group(1) if priority_match else ""
    remainder = priority_match.group(2).strip() if priority_match else group[0]
    effort = infer_effort(group)
    action = group[1] if len(group) >= 2 else ""
    return pad_row([priority, infer_category(remainder), remainder, action, effort], 5)


def split_location_prefix(line: str) -> tuple[str, str]:
    patterns = [
        r"^(Section\s+\d+(?:\.\d+)?)\s+(.+)$",
        r"^(Table\s+\d+(?:\.\d+)?)\s+(.+)$",
        r"^(Figure\s+\d+(?:\.\d+)?)\s+(.+)$",
        r"^(Abstract|Introduction|Methods?|Results?|Discussion|Conclusion|Title)\s+(.+)$",
    ]
    for pattern in patterns:
        match = re.match(pattern, line, re.I)
        if match:
            return match.group(1), match.group(2)
    parts = line.split(" ", 1)
    return parts[0], parts[1] if len(parts) > 1 else ""


def split_table_figure_prefix(line: str) -> tuple[str, str]:
    match = re.match(r"^((?:Table|Figure)\s+\d+(?:\.\d+)?)(?:\s+(.+))?$", line, re.I)
    if match:
        return match.group(1), match.group(2) or ""
    return split_location_prefix(line)


def looks_like_location_start(line: str) -> bool:
    return bool(re.match(r"^(Section\s+\d+(?:\.\d+)?|Table\s+\d+(?:\.\d+)?|Figure\s+\d+(?:\.\d+)?|Abstract|Introduction|Methods?|Results?|Discussion|Conclusion|Title)\b", line, re.I))


def looks_like_table_figure_start(line: str) -> bool:
    return bool(re.match(r"^(Table|Figure)\s+\d+(?:\.\d+)?\b", line, re.I))


def looks_like_priority_start(line: str) -> bool:
    return bool(re.match(r"^(High|Medium|Low)\b", line, re.I))


def infer_issue_type(group: list[str]) -> str:
    joined = " ".join(group).lower()
    for label in ["Grammar", "Clarity", "Wordiness", "Tone", "Repetition", "Tense", "Syntax"]:
        if label.lower() in joined:
            return label
    return "Clarity"


def infer_citation_issue(text: str) -> str:
    lowered = text.lower()
    if "missing" in lowered or "no citation" in lowered:
        return "Missing"
    if "outdated" in lowered:
        return "Outdated"
    if "mismatch" in lowered:
        return "Possibly mismatched"
    return "Weak"


def infer_clarity(group: list[str]) -> str:
    joined = " ".join(group).lower()
    if "unclear" in joined:
        return "Unclear"
    if "partly" in joined:
        return "Partly clear"
    return "Clear" if "clear" in joined else "Partly clear"


def infer_category(text: str) -> str:
    for category in ["Grammar", "Organization", "Factual", "Technical", "Figure", "Citation"]:
        if category.lower() in text.lower():
            return category
    return "Technical"


def infer_effort(group: list[str]) -> str:
    joined = " ".join(group).lower()
    if "high" in joined:
        return "High"
    if "moderate" in joined:
        return "Moderate"
    return "Low"


def split_table_line(line: str, expected_columns: int) -> list[str]:
    if line.strip().startswith("|") and "|" in line.strip()[1:]:
        cells = parse_markdown_row(line)
        if not is_markdown_separator(line):
            return cells

    if "\t" in line:
        return [cell.strip() for cell in line.split("\t") if cell.strip()]

    cells = [cell.strip() for cell in re.split(r"\s{2,}", line) if cell.strip()]
    if len(cells) >= 2:
        return cells

    # Some models collapse tabs to single spaces for the header-example row. Keep
    # those from becoming regular paragraphs by splitting known option-heavy rows.
    if expected_columns == 5 and re.search(r"\b(Grammar|Clarity|Wordiness|Tone|Repetition|Tense|Syntax)\b", line):
        return split_grammar_row(line)
    if expected_columns == 4 and re.search(r"\b(Missing|Weak|Possibly mismatched|Outdated)\b", line):
        return split_citation_row(line)
    if expected_columns == 5 and re.search(r"\b(High|Medium|Low)\b", line) and re.search(r"\b(Low|Moderate|High)\b", line):
        return split_priority_row(line)

    return [line]


def split_grammar_row(line: str) -> list[str]:
    match = re.search(r"\b(Grammar|Clarity|Wordiness|Tone|Repetition|Tense|Syntax)\b", line)
    if not match:
        return [line]
    before = line[: match.start()].strip()
    after = line[match.end() :].strip()
    before_parts = before.split(" ", 1)
    after_parts = after.split(" ", 1)
    return [before_parts[0], before_parts[1] if len(before_parts) > 1 else "", match.group(1), after_parts[0] if after_parts else "", after_parts[1] if len(after_parts) > 1 else ""]


def split_citation_row(line: str) -> list[str]:
    match = re.search(r"\b(Missing|Weak|Possibly mismatched|Outdated)\b", line)
    if not match:
        return [line]
    before = line[: match.start()].strip().split(" ", 1)
    after = line[match.end() :].strip()
    return [before[0], before[1] if len(before) > 1 else "", match.group(1), after]


def split_priority_row(line: str) -> list[str]:
    cells = [cell.strip() for cell in re.split(r"\s{2,}", line) if cell.strip()]
    return cells if len(cells) >= 2 else [line]


def should_stop_table(line: str, rows: list[list[str]]) -> bool:
    if len(rows) <= 1:
        return False
    if is_markdown_separator(line):
        return False
    lowered = line.lower().strip(":")
    if is_numbered_major_heading(line) or is_numbered_subsection_heading(line):
        return True
    if is_bold_heading(line):
        return True
    return any(re.search(pattern, lowered) for pattern in TABLE_STOP_PATTERNS)


def is_bold_heading(line: str) -> bool:
    return bool(re.fullmatch(r"\*\*[^*]{3,}\*\*", line.strip()))


def is_major_heading(line: str) -> bool:
    cleaned = clean_inline_markdown(line)
    return is_numbered_major_heading(cleaned) or cleaned in MAJOR_SECTION_HEADINGS


def is_numbered_major_heading(line: str) -> bool:
    match = re.match(r"^\d+\.\s+(.+)$", clean_inline_markdown(line))
    return bool(match and match.group(1).strip(":") in MAJOR_SECTION_HEADINGS)


def is_numbered_subsection_heading(line: str) -> bool:
    match = re.match(r"^\d+\.\d+\s+(.+)$", clean_inline_markdown(line))
    return bool(match and match.group(1).strip(":") in SUBSECTION_HEADINGS)


def normalize_for_matching(line: str) -> str:
    return re.sub(r"\s+", " ", clean_inline_markdown(line)).strip().lower()


def pad_row(row: list[str], width: int) -> list[str]:
    return row + [""] * (width - len(row))


def normalize_table_rows(rows: list[list[str]]) -> list[list[str]]:
    width = max((len(row) for row in rows), default=0)
    return [row + [""] * (width - len(row)) for row in rows]


def clean_inline_markdown(text: str) -> str:
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"__(.*?)__", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def markdown_to_html(markdown: str) -> str:
    output = []
    list_type = None
    lines = markdown.splitlines()
    index = 0

    def close_list() -> None:
        nonlocal list_type
        if list_type:
            output.append(f"</{list_type}>")
            list_type = None

    def open_list(tag: str) -> None:
        nonlocal list_type
        if list_type == tag:
            return
        close_list()
        output.append(f"<{tag}>")
        list_type = tag

    while index < len(lines):
        known_header = known_table_header(lines[index])
        if known_header:
            close_list()
            rows, index = collect_loose_table(lines, index, known_header)
            output.append(table_to_html(rows))
            continue

        if is_markdown_table(lines, index):
            close_list()
            rows, index = collect_markdown_table(lines, index)
            output.append(table_to_html(rows))
            continue

        if is_tab_table(lines, index):
            close_list()
            rows, index = collect_tab_table(lines, index)
            output.append(table_to_html(rows))
            continue

        raw_line = lines[index]
        line = raw_line.strip()
        if not line:
            close_list()
            index += 1
            continue
        escaped = html.escape(clean_inline_markdown(line))
        if escaped.startswith("# "):
            close_list()
            output.append(f"<h1>{escaped[2:]}</h1>")
        elif escaped.startswith("## "):
            close_list()
            output.append(f"<h2>{escaped[3:]}</h2>")
        elif escaped.startswith("### "):
            close_list()
            output.append(f"<h3>{escaped[4:]}</h3>")
        elif is_bold_heading(line):
            close_list()
            tag = "h1" if is_major_heading(line) else "h2"
            output.append(f"<{tag}>{html.escape(clean_inline_markdown(line))}</{tag}>")
        elif re.match(r"^[-*]\s+", escaped):
            open_list("ul")
            item = re.sub(r"^[-*]\s+", "", escaped)
            output.append(f"<li>{item}</li>")
        elif re.match(r"^\d+\.\s+", escaped):
            open_list("ol")
            item = re.sub(r"^\d+\.\s+", "", escaped)
            output.append(f"<li>{item}</li>")
        else:
            close_list()
            output.append(f"<p>{escaped}</p>")
        index += 1
    close_list()
    return "".join(output)


def table_to_html(rows: list[list[str]]) -> str:
    rows = normalize_table_rows(rows)
    if not rows:
        return ""

    header = "".join(f"<th>{html.escape(clean_inline_markdown(cell))}</th>" for cell in rows[0])
    body_rows = []
    for row in rows[1:]:
        cells = "".join(f"<td>{html.escape(clean_inline_markdown(cell))}</td>" for cell in row)
        body_rows.append(f"<tr>{cells}</tr>")
    return f"<table><thead><tr>{header}</tr></thead><tbody>{''.join(body_rows)}</tbody></table>"


def normalize_markdown(markdown: str) -> str:
    return re.sub(r"\n{3,}", "\n\n", markdown.strip())
